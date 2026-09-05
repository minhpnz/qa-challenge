"""Generate the test case workbook and import-ready CSVs from test_case_data.py."""

import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_case_data import CART, DEFECTS, LOGIN  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_XLSX = os.path.join(ROOT, "docs", "DemoBlaze-Test-Cases.xlsx")
OUT_CSV_DIR = os.path.join(ROOT, "docs", "test-cases-csv")

NAVY = "1F3864"
LIGHT = "D9E2F3"
BAND = "F2F5FB"
AMBER = "FFF2CC"
GREY = "808080"

# Design columns (what the case is) then execution columns (what happened when it
# was run) then traceability. The execution block ships empty on purpose: this is
# a test case suite, and a suite that arrives pre-filled with results nobody
# produced is a fiction. A tester fills Actual Result / Status / Executed on a
# real cycle; automated cases report through CI instead — see the Summary tab.
HEADERS = [
    "Test Case ID", "Module", "Feature Area", "Title", "Type", "Priority",
    "Preconditions", "Test Steps", "Test Data", "Expected Result",
    "Actual Result", "Status", "Executed By", "Executed On", "Environment",
    "Automated", "Automation Reference", "Defect ID",
]
WIDTHS = [14, 10, 18, 46, 13, 9, 30, 46, 26, 58, 34, 12, 14, 14, 20, 11, 52, 12]

DEFECT_HEADERS = [
    "Defect ID", "Area", "Severity", "Likelihood", "Status", "Summary",
    "Detail", "Test Case ID", "Pinned By",
]
DEFECT_WIDTHS = [11, 12, 10, 11, 9, 46, 72, 20, 60]

# Light verticals to separate columns, a darker rule under every row. Tall
# wrapped cells swallow a uniformly light grid, which is what made the earlier
# version hard to scan.
VLINE = Side(style="thin", color="D6DCE4")
HLINE = Side(style="thin", color="8496B0")
BORDER = Border(left=VLINE, right=VLINE, top=None, bottom=HLINE)
HEAD_BORDER = Border(left=VLINE, right=VLINE, top=None, bottom=Side(style="medium", color=NAVY))


def estimate_height(values, widths):
    """Approximate the wrapped height so rows are readable without relying on the
    viewer's auto-fit, which Excel and Sheets apply inconsistently to wrapped text."""
    lines = 1
    for value, width in zip(values, widths):
        text = str(value)
        wrapped = sum(max(1, -(-len(seg) // max(8, int(width * 0.95)))) for seg in text.split("\n"))
        lines = max(lines, wrapped)
    return min(14 * min(lines, 12) + 6, 180)


def style_header(ws, headers, widths):
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", start_color=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = HEAD_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_cases(ws, rows, module):
    style_header(ws, HEADERS, WIDTHS)
    for i, (cid, area, title, ttype, prio, pre, steps, data, expected, ref, defect) in enumerate(rows):
        r = i + 2
        automated = "No" if not ref or ref == "Manual" else "Yes"
        values = [cid, module, area, title, ttype, prio, pre, steps, data, expected,
                  "", "Not Run", "", "", "", automated, ref or "—", defect or "—"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if i % 2:
                cell.fill = PatternFill("solid", start_color=BAND)
        ws.cell(row=r, column=1).font = Font(bold=True)
        if prio == "P0":
            ws.cell(row=r, column=6).font = Font(bold=True, color="C00000")
        ws.cell(row=r, column=12).font = Font(color=GREY, italic=True)
        ws.cell(row=r, column=12).alignment = Alignment(horizontal="center", vertical="top")
        # Editable cells get a paler ground so a tester can see where to type.
        for col in (11, 12, 13, 14, 15):
            ws.cell(row=r, column=col).fill = PatternFill("solid", start_color="FFFDF5")
        ws.cell(row=r, column=16).font = (
            Font(color="1F7A1F", bold=True) if automated == "Yes" else Font(color=GREY)
        )
        if defect:
            ws.cell(row=r, column=18).fill = PatternFill("solid", start_color=AMBER)
        ws.row_dimensions[r].height = estimate_height(values, WIDTHS)

    last = len(rows) + 1
    for column, options in (
        ("E", '"Functional,Negative,Edge case,Security,UI"'),
        ("F", '"P0,P1,P2"'),
        ("L", '"Not Run,Pass,Fail,Blocked,N/A"'),
        ("P", '"Yes,No"'),
    ):
        dv = DataValidation(type="list", formula1=options, allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"{column}2:{column}{last}")
    return last


def write_defects(ws, rows):
    style_header(ws, DEFECT_HEADERS, DEFECT_WIDTHS)
    severity_colour = {"High": "C00000", "Medium": "BF8F00", "Low": "808080"}
    for i, row in enumerate(rows):
        r = i + 2
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if i % 2:
                cell.fill = PatternFill("solid", start_color=BAND)
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=3).font = Font(bold=True, color=severity_colour.get(row[2], "000000"))


def write_summary(ws, n_login, n_cart, n_defects):
    """Summary is laid out as three bordered tables.

    Gridlines are off, so every block draws its own frame — text floating on a
    blank sheet with no rules is unreadable, which is exactly what an earlier
    version of this tab did.
    """
    for col, width in zip("ABCDE", (44, 12, 12, 12, 78)):
        ws.column_dimensions[col].width = width

    def table_header(row, labels, span_last=False):
        for offset, text in enumerate(labels):
            c = ws.cell(row=row, column=1 + offset, value=text)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", start_color=NAVY)
            c.border = Border(left=VLINE, right=VLINE, bottom=Side(style="medium", color=NAVY))
            c.alignment = Alignment(horizontal="left" if offset == 0 else "center", vertical="center")
        if span_last:
            for col in range(len(labels) + 1, 6):
                c = ws.cell(row=row, column=col)
                c.fill = PatternFill("solid", start_color=NAVY)
                c.border = Border(bottom=Side(style="medium", color=NAVY))
        ws.row_dimensions[row].height = 20

    def frame(row, last_col=5, band=False, bold=False, top_rule=False):
        for col in range(1, last_col + 1):
            c = ws.cell(row=row, column=col)
            c.border = Border(
                left=VLINE, right=VLINE,
                top=Side(style="medium", color=NAVY) if top_rule else None,
                bottom=HLINE,
            )
            if band:
                c.fill = PatternFill("solid", start_color=BAND)
            if bold:
                c.font = Font(bold=True)
            if col in (2, 3, 4):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(vertical="center", wrap_text=True)

    title = ws.cell(row=1, column=1, value="DemoBlaze — Login & Cart Test Case Suite")
    title.font = Font(bold=True, size=16, color=NAVY)
    ws.cell(row=2, column=1, value="Application under test: https://www.demoblaze.com/").font = Font(italic=True)
    ws.cell(row=3, column=1, value="Scope: Login (incl. registration as a dependency) and Cart / Checkout")
    ws.cell(row=4, column=1,
            value="Generated from scripts/test_case_data.py — regenerate with: python3 scripts/generate_test_cases.py"
            ).font = Font(italic=True, color=GREY)

    # ---------------- coverage ----------------
    ws.cell(row=6, column=1, value="Coverage").font = Font(bold=True, size=13, color=NAVY)
    table_header(7, ["Metric", "Login", "Cart", "Total", "Notes"])

    login_last, cart_last = n_login + 1, n_cart + 1

    def counts(row, name, col_letter, criterion, note="", band=False, bold=False):
        ws.cell(row=row, column=1, value=name)
        if criterion is None:
            ws.cell(row=row, column=2, value=f"=COUNTA('Login Test Cases'!$A$2:$A${login_last})")
            ws.cell(row=row, column=3, value=f"=COUNTA('Cart Test Cases'!$A$2:$A${cart_last})")
        else:
            ws.cell(row=row, column=2,
                    value=f"=COUNTIF('Login Test Cases'!${col_letter}$2:${col_letter}${login_last},\"{criterion}\")")
            ws.cell(row=row, column=3,
                    value=f"=COUNTIF('Cart Test Cases'!${col_letter}$2:${col_letter}${cart_last},\"{criterion}\")")
        ws.cell(row=row, column=4, value=f"=B{row}+C{row}")
        ws.cell(row=row, column=5, value=note)
        frame(row, band=band, bold=bold)

    counts(8, "Total test cases", None, None, "44 Login + 50 Cart.", bold=True)
    for offset, ttype in enumerate(["Functional", "Negative", "Edge case", "Security", "UI"]):
        counts(9 + offset, f"    {ttype}", "E", ttype, band=(offset % 2 == 0))
    counts(14, "P0 — blocks release", "F", "P0",
           "Login, add to cart, checkout, cross-account cart isolation.", bold=True)
    counts(15, "Automated", "P", "Yes", "Covered by the Playwright suite in this repository.", band=True)
    counts(16, "Manual / exploratory", "P", "No",
           "Deliberately not automated: browser chrome (Escape, backdrop, browser Back), "
           "network-fault injection, and security probes needing a controlled environment.")

    ws.cell(row=17, column=1, value="Automation coverage").font = Font(bold=True)
    pct = ws.cell(row=17, column=4, value="=IF(D8=0,0,D15/D8)")
    pct.number_format = "0.0%"
    ws.cell(row=17, column=5,
            value="Share of documented cases with an executable check. P0 coverage is what gates a "
                  "release, not this number.")
    frame(17, band=True, bold=False)
    ws.cell(row=17, column=1).font = Font(bold=True)

    # ---------------- defects ----------------
    ws.cell(row=19, column=1, value="Defects raised").font = Font(bold=True, size=13, color=NAVY)
    table_header(20, ["Severity", "Count", "", "", "Notes"])
    ws.cell(row=21, column=1, value="Total logged")
    ws.cell(row=21, column=2, value=f"=COUNTA(Defects!$A$2:$A${n_defects + 1})")
    ws.cell(row=21, column=5, value="Full detail, with verification method, on the Defects tab.")
    frame(21, bold=True)
    severity_note = {
        "High": "Empty-cart checkout, no login rate limiting, unmasked card number, "
                "confirmation shown before the server confirms, HTTP 500 on empty username.",
        "Medium": "Receipt dated a month early, presence-only order validation, username "
                  "enumeration, no password policy, silent login failure when the API is down.",
        "Low": "Anonymous cart discarded on login, duplicate DOM ids, trailing newline in a "
               "product title, inconsistent confirmation copy, Enter does not submit.",
    }
    for offset, severity in enumerate(["High", "Medium", "Low"]):
        row = 22 + offset
        ws.cell(row=row, column=1, value=f"    {severity} severity")
        ws.cell(row=row, column=2, value=f"=COUNTIF(Defects!$C$2:$C${n_defects + 1},\"{severity}\")")
        ws.cell(row=row, column=5, value=severity_note[severity])
        frame(row, band=(offset % 2 == 0))
        ws.cell(row=row, column=2).font = Font(bold=True, color={"High": "C00000", "Medium": "BF8F00", "Low": GREY}[severity])

    # ---------------- legend ----------------
    ws.cell(row=27, column=1, value="How to read this workbook").font = Font(bold=True, size=13, color=NAVY)
    table_header(28, ["Column / tab", "Meaning"])
    ws.merge_cells(start_row=28, start_column=2, end_row=28, end_column=5)
    notes = [
        ("Login Test Cases", "44 cases: authentication, session handling, validation, security, boundaries."),
        ("Cart Test Cases", "50 cases: add-to-cart, cart management, persistence, isolation, checkout."),
        ("Defects", "Every issue found in the application, cross-referenced to the case that exposes it."),
        ("Type", "Functional = intended behaviour. Negative = invalid input rejected. Edge case = boundary "
                 "or unusual-but-valid. Security / UI as labelled."),
        ("Priority", "P0 = blocks release. P1 = important. P2 = lower risk."),
        ("Automated", "Yes = an executable check exists. No = deliberately manual, for the reasons above."),
        ("Actual Result / Status", "Execution columns, shipped empty and shaded. A tester fills them during a "
                                   "run; Status defaults to Not Run and is a dropdown (Pass / Fail / Blocked / N/A). "
                                   "Automated cases report through CI instead of here — a suite delivered with "
                                   "results nobody produced is a fiction."),
        ("Executed By / On / Environment", "Who ran the cycle, when, and against which browser and build, so a "
                                            "result can be reproduced rather than taken on trust."),
        ("Automation Reference", "The exact spec and test title that runs the case, so a reviewer can go "
                                 "from a row straight to running code."),
        ("Expected Result", "Where the app is defective the row states BOTH the expected and the observed "
                            "behaviour and names the defect. A case that records a bug as correct behaviour "
                            "launders the defect into a requirement."),
    ]
    for offset, (name, text) in enumerate(notes):
        row = 29 + offset
        ws.cell(row=row, column=1, value=name).font = Font(bold=True)
        ws.cell(row=row, column=2, value=text)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        frame(row, band=(offset % 2 == 0))
        ws.cell(row=row, column=2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30 if len(text) > 96 else 18

    ws.freeze_panes = "A5"


def dump_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def main():
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    os.makedirs(OUT_CSV_DIR, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    n_login = write_cases(wb.create_sheet("Login Test Cases"), LOGIN, "Login") - 1
    n_cart = write_cases(wb.create_sheet("Cart Test Cases"), CART, "Cart") - 1
    write_defects(wb.create_sheet("Defects"), DEFECTS)
    write_summary(summary, n_login, n_cart, len(DEFECTS))
    summary.sheet_view.showGridLines = False
    wb.save(OUT_XLSX)

    def to_rows(cases, module):
        out = []
        for cid, area, title, ttype, prio, pre, steps, data, expected, ref, defect in cases:
            automated = "No" if not ref or ref == "Manual" else "Yes"
            out.append([cid, module, area, title, ttype, prio, pre, steps, data, expected,
                        "", "Not Run", "", "", "", automated, ref or "—", defect or "—"])
        return out

    dump_csv(os.path.join(OUT_CSV_DIR, "01-summary.csv"),
             ["Section", "Value"],
             [["Application under test", "https://www.demoblaze.com/"],
              ["Scope", "Login (incl. registration dependency) and Cart / Checkout"],
              ["Total test cases", len(LOGIN) + len(CART)],
              ["Login test cases", len(LOGIN)],
              ["Cart test cases", len(CART)],
              ["Automated", sum(1 for r in LOGIN + CART if r[9] and r[9] != "Manual")],
              ["Manual / exploratory", sum(1 for r in LOGIN + CART if not r[9] or r[9] == "Manual")],
              ["Defects raised", len(DEFECTS)]])
    dump_csv(os.path.join(OUT_CSV_DIR, "02-login-test-cases.csv"), HEADERS, to_rows(LOGIN, "Login"))
    dump_csv(os.path.join(OUT_CSV_DIR, "03-cart-test-cases.csv"), HEADERS, to_rows(CART, "Cart"))
    dump_csv(os.path.join(OUT_CSV_DIR, "04-defects.csv"), DEFECT_HEADERS, [list(r) for r in DEFECTS])

    print(f"xlsx: {OUT_XLSX}")
    print(f"csv:  {OUT_CSV_DIR}")
    print(f"login={len(LOGIN)} cart={len(CART)} defects={len(DEFECTS)}")


if __name__ == "__main__":
    main()
